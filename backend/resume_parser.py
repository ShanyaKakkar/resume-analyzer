from pathlib import Path

import docx
import pdfplumber


def extract_text_from_pdf(path):
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n".join(pages)


def extract_text_from_docx(path):
    document = docx.Document(path)
    lines = []

    # The page header often holds the name and contact details
    for p in document.sections[0].header.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())

    # Normal paragraphs
    for p in document.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())

    # Tables (many resume templates use tables for the layout)
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in seen:
                    seen.add(cell_text)
                    lines.append(cell_text)

    return "\n".join(lines)


def extract_text(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".docx":
        return extract_text_from_docx(path)
    raise ValueError("Only PDF and DOCX files are supported")